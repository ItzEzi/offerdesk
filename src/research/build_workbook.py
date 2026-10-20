import datetime as dt, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as CL
from openpyxl.comments import Comment

OUT = sys.argv[1] if len(sys.argv) > 1 else '/mnt/user-data/outputs/Market Research Workbook.xlsx'
wb = Workbook()
try:
    wb._named_styles['Normal'].font = Font(name='Arial', size=10)
except Exception:
    pass

ARIAL = 'Arial'
NAVY = '1F3864'
F_IN = PatternFill('solid', fgColor='FFF2CC')        # light yellow: you type here
F_HEAD = PatternFill('solid', fgColor=NAVY)
F_GROUP = PatternFill('solid', fgColor='D9E1F2')
F_CALC = PatternFill('solid', fgColor='F2F2F2')
F_KEY = PatternFill('solid', fgColor='E2EFDA')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def font(color='000000', bold=False, italic=False, size=10, underline=None):
    return Font(name=ARIAL, color=color, bold=bold, italic=italic, size=size, underline=underline)
BLUE, BLACK, GREEN = '0000FF', '000000', '008000'

FMT = {'$': '$#,##0;($#,##0);-', '%': '0.0%;-0.0%;-', 'n': '#,##0;-#,##0;-', 'n1': '0.0;-0.0;-', 'd': 'yyyy-mm-dd', 't': '@', 'yr': '0'}

def title(ws, text, sub):
    ws['A1'] = text; ws['A1'].font = font(NAVY, True, size=14)
    ws['A2'] = sub; ws['A2'].font = font('595959', italic=True)
    ws.row_dimensions[1].height = 22

def header(ws, row, cols):
    for i, (name, width, *_r) in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = font('FFFFFF', True); c.fill = F_HEAD; c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        ws.column_dimensions[CL(i)].width = width
    ws.row_dimensions[row].height = 42

def groups(ws, row, spans):
    for (a, b, text) in spans:
        ws.merge_cells(start_row=row, start_column=a, end_row=row, end_column=b)
        c = ws.cell(row=row, column=a, value=text)
        c.font = font(NAVY, True); c.fill = F_GROUP; c.alignment = Alignment(horizontal='center')
        for k in range(a, b + 1):
            ws.cell(row=row, column=k).fill = F_GROUP; ws.cell(row=row, column=k).border = BOX

def table(ws, first, last, cols, example=None, formulas=None):
    """cols: (header, width, kind, fmt). kind: in | calc | link(green). formulas: {col_index: lambda r: '=...'}"""
    formulas = formulas or {}
    for r in range(first, last + 1):
        for i, (name, width, kind, fmt) in enumerate(cols, 1):
            c = ws.cell(row=r, column=i)
            c.border = BOX
            if fmt:
                c.number_format = FMT.get(fmt, fmt)
            if kind == 'in':
                c.fill = F_IN; c.font = font(BLUE)
            else:
                c.fill = F_CALC; c.font = font(GREEN if kind == 'link' else BLACK)
                if i in formulas:
                    c.value = formulas[i](r)
    if example:
        for i, v in example.items():
            c = ws.cell(row=first, column=i, value=v)
            c.font = font(BLUE, italic=True)

def note(ws, cell, text):
    ws[cell].comment = Comment(text, 'Workbook')

# ============================ Start Here ============================
ws = wb.active; ws.title = 'Start Here'
ws.sheet_view.showGridLines = False
ws.column_dimensions['A'].width = 4
ws.column_dimensions['B'].width = 34
ws.column_dimensions['C'].width = 62
ws.column_dimensions['D'].width = 58
ws['B1'] = 'Market Research Workbook'; ws['B1'].font = font(NAVY, True, size=16)
ws['B2'] = 'Your own research, your own numbers. Every calculation is a plain formula you can click on and check.'
ws['B2'].font = font('595959', italic=True)
r = 4
ws.cell(row=r, column=2, value='How the colors work').font = font(NAVY, True, size=12); r += 1
leg = [('Yellow cell, blue text', 'You type it (from a source you checked).', F_IN, BLUE),
       ('Gray cell, black text', 'Calculated. Don\'t type here; click it to see the formula.', F_CALC, BLACK),
       ('Gray cell, green text', 'Pulled in from another tab.', F_CALC, GREEN),
       ('Light blue cell', 'Filled in by OfferDesk from public data (Research pack). Spot-check a few against the source.', PatternFill('solid', fgColor='DDEBF7'), BLUE),
       ('Italic first row', 'An EXAMPLE row ("Sampleville, OH") so you can see the format. Delete it once you have your own rows.', F_IN, BLUE)]
for a, b, fill, col in leg:
    c = ws.cell(row=r, column=2, value=a); c.fill = fill; c.font = font(col); c.border = BOX
    ws.cell(row=r, column=3, value=b).font = font()
    r += 1
r += 1
ws.cell(row=r, column=2, value='The plan (about 4–5 weeks for ALL your metros, done side by side)').font = font(NAVY, True, size=12); r += 1
steps = [
    ('1. Metros', 'Week 1: write your 3–5 candidate metros in the Metros tab (column A), same name everywhere. Do steps 2–4 for all of them, then cut to your best 2.'),
    ('2. ZIPs', 'For each metro, add 5–8 ZIPs. Faster: star the metros in OfferDesk → Markets → Research pack, and it fills this tab (light blue). Then spot-check 3 ZIPs per metro with the Redfin link.'),
    ('3. Listing Watch', 'Save 20–30 listings in your price range per metro. Update status and price once a week.'),
    ('4. Past Flips', 'Find houses that sold cheap, then resold renovated within 12 months. OfferDesk lists likely ones in Flip Candidates; you check each one\'s sale history and type the earlier sale.'),
    ('5. Calls', 'Weeks 2–3, your 2 finalists only: call 10 locals each (3 agents, 2 property managers, 2 contractors, 1 hard-money lender, 1 title company). Same questions every time.'),
    ('6. Metros', 'Everything rolls up here. Set your own weights; the score ranks your metros against each other.'),
    ('7. Paper Flips', 'Run your numbers on 10 real listings in the Deal Check tab, log them here, then check what they really sold for.'),
    ('8. Ready check', 'Bottom of the Metros tab. Don\'t buy in a metro until every box says Yes.'),
    ('Weekly', 'In OfferDesk, drop your new Redfin downloads, then "Update my workbook": it refreshes the light-blue rows and adds new listings. Rows you edit yourself are never overwritten once you change their Source/Notes cell.'),
]
for a, b in steps:
    ws.cell(row=r, column=2, value=a).font = font(bold=True)
    c = ws.cell(row=r, column=3, value=b); c.font = font(); c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    ws.row_dimensions[r].height = 28
    r += 1
r += 1
ws.cell(row=r, column=2, value='Where to get each number').font = font(NAVY, True, size=12); r += 1
for i, h in enumerate(['Number', 'Where (free)', 'How'], 2):
    c = ws.cell(row=r, column=i, value=h); c.font = font('FFFFFF', True); c.fill = F_HEAD; c.border = BOX
r += 1
src = [
    ('Median sale price, homes sold, days on market, sale-to-list, % sold above list, price change, active listings',
     'Redfin ZIP housing market page', 'https://www.redfin.com/zipcode/44102/housing-market',
     'Swap in your ZIP (the ZIPs tab builds this link for you). Use the "last 3 months" numbers and stay consistent.'),
    ('Same numbers, all ZIPs at once', 'Redfin Data Center', 'https://www.redfin.com/news/data-center/',
     'Download "ZIP code" market tracker data if you want to check many ZIPs.'),
    ('Typical rent, 3-bed', 'HUD Fair Market Rents (by ZIP), or Zillow rentals', 'https://www.huduser.gov/portal/datasets/fmr.html',
     'Use the Small Area FMR for the ZIP. Cross-check with 3–5 real rentals on Zillow, and ask a property manager.'),
    ('Owner vs. rented single-family homes', 'Census table B25032 (data.census.gov)', 'https://data.census.gov/table/ACSDT5Y2024.B25032',
     'Line 3 = owner-occupied "1, detached". Line 14 = renter-occupied "1, detached". The ZIPs tab link opens it for your ZIP; if it doesn\'t, pick Geography → ZIP Code Tabulation Area and type the ZIP.'),
    ('Median year built', 'Census table B25035', 'https://data.census.gov/table/ACSDT5Y2024.B25035', 'One number per ZIP.'),
    ('Median household income', 'Census table B19013, or Census Reporter', 'https://censusreporter.org/',
     'Census Reporter has an easy page per ZIP (the ZIPs tab links it).'),
    ('County jobs (2-year change)', 'BLS QCEW (county employment)', 'https://www.bls.gov/cew/',
     'Compare total employment now vs. the same quarter 2 years ago.'),
    ('County population since 2020', 'Census population estimates', 'https://www.census.gov/programs-surveys/popest.html',
     'County totals. Compare the latest year to 2020.'),
    ('Past flips', 'Zillow / Redfin "sale history", county recorder or assessor site', '',
     'A house that sold, then sold again within 12 months for a lot more, is almost always a flip.'),
    ('Flip rates for big metros', 'ATTOM home flipping report (quarterly)', 'https://www.attomdata.com/news/market-trends/flipping/',
     'Flips as a % of all sales. The US average was 8% in early 2026.'),
    ('"We buy houses" competition', 'Google, Facebook groups, the local investor club (REIA)', '',
     'Search "we buy houses [city]" and count the different companies on the first 3 pages, ads included.'),
]
for a, where, url, how in src:
    ws.cell(row=r, column=2, value=a).font = font()
    c = ws.cell(row=r, column=3, value=where if not url else f'{where}\n{url}')
    c.font = font('0563C1', underline='single') if url else font()
    if url:
        c.hyperlink = url
    ws.cell(row=r, column=4, value=how).font = font()
    for k in (2, 3, 4):
        ws.cell(row=r, column=k).alignment = Alignment(wrap_text=True, vertical='top'); ws.cell(row=r, column=k).border = BOX
    ws.row_dimensions[r].height = 44
    r += 1
r += 1
ws.cell(row=r, column=2, value='Ask every local the same questions').font = font(NAVY, True, size=12); r += 1
qs = ['What does a renovated 3-bed, 1–2 bath sell for in [ZIP] right now?',
      'What does a cosmetic rehab cost per square foot here? A full gut?',
      'What would that house rent for?',
      'Which areas sell fast, and which ones sit? Where would you not buy?',
      'How many investors and flippers are buying here? (1 = few, 5 = crowded)',
      'Who are the active flippers? Which lenders and contractors do they use?',
      'Title companies: what share of your closings are off-market or wholesale deals?',
      'Would you work with an out-of-state buyer? What do you need from me?',
      'Who else should I talk to?']
for q in qs:
    c = ws.cell(row=r, column=2, value='•  ' + q); c.font = font()
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    r += 1
r += 1
ws.cell(row=r, column=2, value='Rules').font = font(NAVY, True, size=12); r += 1
for t in ['Any number that decides money needs two independent sources, and at least one should be a local person.',
          'A sold price is a fact. Anyone\'s opinion (an agent\'s, a course\'s, an AI\'s) is a guess until a sold price backs it.',
          'Use the lowest comp for your resale price, add 20% to rehab, and the deal must still work in the stress test (Deal Check tab).']:
    c = ws.cell(row=r, column=2, value='•  ' + t); c.font = font()
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    r += 1

# ============================ ZIPs ============================
N_Z = 100; Z0 = 8; Z1 = Z0 + N_Z - 1
zs = wb.create_sheet('ZIPs')
title(zs, 'ZIPs', 'One row per ZIP. Click the links on the right, copy the numbers into the yellow cells. Keep "Metro" spelled exactly as in the Metros tab.')
zcols = [
    ('Metro', 20, 'in', 't'), ('ZIP', 8, 'in', 't'), ('City', 14, 'in', 't'),
    ('Median sale price ($)', 12, 'in', '$'), ('Homes sold, last 3 mo', 9, 'in', 'n'), ('Active listings now', 9, 'in', 'n'),
    ('Median days on market', 9, 'in', 'n'), ('Sale-to-list', 8, 'in', '%'), ('Sold above list', 8, 'in', '%'), ('Price change, 1 yr', 8, 'in', '%'),
    ('Typical rent, 3-bed ($/mo)', 10, 'in', '$'),
    ('Median year built', 8, 'in', 'yr'), ('Owner-occupied single-family homes (B25032 line 3)', 12, 'in', 'n'), ('Rented single-family homes (B25032 line 14)', 12, 'in', 'n'),
    ('Median household income ($)', 11, 'in', '$'),
    ('Checked on', 11, 'in', 'd'), ('Source / notes', 24, 'in', 't'),
    ('Months of supply', 8, 'calc', 'n1'), ('Rent yield (gross)', 8, 'calc', '%'), ('Single-family homes rented out', 9, 'calc', '%'),
    ('Flips found, 12 mo', 8, 'link', 'n'), ('Flips as % of sales (est.)', 9, 'calc', '%'), ('Listings watched', 8, 'link', 'n'),
    ('Redfin', 9, 'calc', None), ('Census B25032', 10, 'calc', None), ('Census Reporter', 10, 'calc', None)]
groups(zs, Z0 - 2, [(1, 3, 'Where'), (4, 10, 'From Redfin (ZIP page)'), (11, 11, 'Rent'), (12, 15, 'From the Census'), (16, 17, 'Your check'),
                    (18, 23, 'Calculated'), (24, 26, 'Open the source')])
header(zs, Z0 - 1, zcols)
PF = "'Past Flips'"; LW = "'Listing Watch'"; FC = "'Flip Candidates'"
zf = {
    18: lambda r: f'=IF(OR(E{r}="",F{r}="",E{r}=0),"",F{r}/(E{r}/3))',
    19: lambda r: f'=IF(OR(D{r}="",K{r}="",D{r}=0),"",K{r}*12/D{r})',
    20: lambda r: f'=IF(OR(M{r}="",N{r}=""),"",IF(M{r}+N{r}=0,"",N{r}/(M{r}+N{r})))',
    21: lambda r: f'=IF(B{r}="","",COUNTIF({PF}!$B$8:$B$307,B{r})+COUNTIFS({FC}!$B$8:$B$207,B{r},{FC}!$O$8:$O$207,"Likely flip"))',
    22: lambda r: f'=IF(OR(U{r}="",E{r}="",E{r}=0),"",U{r}/(E{r}*4))',
    23: lambda r: f'=IF(B{r}="","",COUNTIF({LW}!$B$8:$B$307,B{r}))',
    24: lambda r: f'=IF(B{r}="","",HYPERLINK("https://www.redfin.com/zipcode/"&B{r}&"/housing-market","Redfin"))',
    25: lambda r: f'=IF(B{r}="","",HYPERLINK("https://data.census.gov/table/ACSDT5Y2024.B25032?g=860XX00US"&B{r},"B25032"))',
    26: lambda r: f'=IF(B{r}="","",HYPERLINK("https://censusreporter.org/profiles/86000US"&B{r}&"-"&B{r}&"/","Profile"))',
}
table(zs, Z0, Z1, zcols, formulas=zf, example={
    1: 'Sampleville, OH', 2: 'EX001', 3: 'Sampleville', 4: 185000, 5: 42, 6: 38, 7: 34, 8: 0.975, 9: 0.22, 10: 0.041, 11: 1450,
    12: 1958, 13: 2410, 14: 520, 15: 61500, 16: dt.date(2026, 10, 15), 17: 'EXAMPLE row: delete me'})
for c in (24, 25, 26):
    for r in range(Z0, Z1 + 1):
        zs.cell(row=r, column=c).font = font('0563C1', underline='single')
zs.freeze_panes = zs.cell(row=Z0, column=3)
note(zs, 'H7', 'Average sale price ÷ list price. 0.98 = homes sell for 98% of asking. Type 98% or 0.98.')
note(zs, 'R7', 'Active listings ÷ monthly sales. Under 3 = homes go fast (sellers\' market). Over 6 = slow (buyers\' market).')
note(zs, 'T7', 'Rented single-family ÷ all occupied single-family homes. Most rented houses are owned by investors, so a high or rising share means more investors already here.')
note(zs, 'V7', 'Flips you found in the last 12 months ÷ (3-month sales × 4). A rough flip rate. The US average is about 8%. Only meaningful once you have searched the sales history properly.')

# ============================ Listing Watch ============================
N_L = 300; L0 = 8; L1 = L0 + N_L - 1
ls = wb.create_sheet('Listing Watch')
title(ls, 'Listing Watch', 'Save 20–30 listings per metro in your price range. Once a week, update Current price, Status and dates. This shows how fast deals really go.')
lcols = [('Metro', 18, 'in', 't'), ('ZIP', 8, 'in', 't'), ('Address', 26, 'in', 't'), ('Link', 16, 'in', 't'), ('Beds', 6, 'in', 'n'), ('Baths', 6, 'in', 'n1'),
         ('Sqft', 8, 'in', 'n'), ('Condition', 11, 'in', 't'), ('Listed on', 11, 'in', 'd'), ('List price ($)', 11, 'in', '$'), ('Current price ($)', 11, 'in', '$'),
         ('Status', 11, 'in', 't'), ('Went pending on', 11, 'in', 'd'), ('Sold on', 11, 'in', 'd'), ('Sold price ($)', 11, 'in', '$'), ('Last checked', 11, 'in', 'd'),
         ('Notes', 24, 'in', 't'),
         ('Price cut', 8, 'calc', '%'), ('Days to pending', 8, 'calc', 'n'), ('Sold vs. first list price', 9, 'calc', '%'), ('List $/sqft', 8, 'calc', '$'), ('Sold $/sqft', 8, 'calc', '$')]
groups(ls, L0 - 2, [(1, 8, 'The house'), (9, 17, 'Track it weekly'), (18, 22, 'Calculated')])
header(ls, L0 - 1, lcols)
lf = {18: lambda r: f'=IF(OR(J{r}="",K{r}="",J{r}=0),"",1-K{r}/J{r})',
      19: lambda r: f'=IF(OR(I{r}="",M{r}=""),"",M{r}-I{r})',
      20: lambda r: f'=IF(OR(J{r}="",O{r}="",J{r}=0),"",O{r}/J{r})',
      21: lambda r: f'=IF(OR(J{r}="",G{r}="",G{r}=0),"",J{r}/G{r})',
      22: lambda r: f'=IF(OR(O{r}="",G{r}="",G{r}=0),"",O{r}/G{r})'}
table(ls, L0, L1, lcols, formulas=lf, example={1: 'Sampleville, OH', 2: 'EX001', 3: '123 Example St', 5: 3, 6: 1, 7: 1150, 8: 'Fixer', 9: dt.date(2026, 9, 1),
                                               10: 149900, 11: 139900, 12: 'Pending', 13: dt.date(2026, 10, 2), 16: dt.date(2026, 10, 15), 17: 'EXAMPLE row: delete me'})
ls.freeze_panes = ls.cell(row=L0, column=4)
dv = DataValidation(type='list', formula1='"Fixer,Dated,Renovated,New"', allow_blank=True); ls.add_data_validation(dv); dv.add(f'H{L0}:H{L1}')
dv = DataValidation(type='list', formula1='"Active,Pending,Sold,Off market"', allow_blank=True); ls.add_data_validation(dv); dv.add(f'L{L0}:L{L1}')

# ============================ Past Flips ============================
N_P = 300; P0 = 8; P1 = P0 + N_P - 1
ps = wb.create_sheet('Past Flips')
title(ps, 'Past Flips', 'Houses that sold, then resold for a lot more within ~12 months. They prove flips work here, show the profit, and show how many flippers are active.')
ps['A3'] = 'Rehab $/sqft to assume when you don\'t know it'; ps['A3'].font = font(bold=True)
ps['E3'] = 30; ps['E3'].fill = F_IN; ps['E3'].font = font(BLUE); ps['E3'].number_format = '$#,##0'; ps['E3'].border = BOX
ps['F3'] = 'Your assumption. Replace it with what local contractors tell you (Calls tab).'; ps['F3'].font = font('595959', italic=True)
ps['A4'] = 'Buy + sell + holding costs, as % of the resale price'; ps['A4'].font = font(bold=True)
ps['E4'] = 0.10; ps['E4'].fill = F_IN; ps['E4'].font = font(BLUE); ps['E4'].number_format = '0.0%'; ps['E4'].border = BOX
ps['F4'] = 'Your assumption: agent fees, closing costs, taxes, insurance, utilities while holding. 10% is a common rough figure.'; ps['F4'].font = font('595959', italic=True)
pcols = [('Metro', 18, 'in', 't'), ('ZIP', 8, 'in', 't'), ('Address', 26, 'in', 't'), ('Sqft', 8, 'in', 'n'), ('Bought on', 11, 'in', 'd'), ('Bought for ($)', 11, 'in', '$'),
         ('Sold on', 11, 'in', 'd'), ('Sold for ($)', 11, 'in', '$'), ('Rehab ($) if known', 11, 'in', '$'), ('Flipper (name or LLC)', 18, 'in', 't'), ('Where you found it', 14, 'in', 't'),
         ('Notes', 22, 'in', 't'),
         ('Months held', 7, 'calc', 'n1'), ('Price jump ($)', 10, 'calc', '$'), ('Bought $/sqft', 8, 'calc', '$'), ('Sold $/sqft', 8, 'calc', '$'),
         ('Rehab used ($)', 10, 'calc', '$'), ('Other costs ($)', 10, 'calc', '$'), ('Est. profit ($)', 11, 'calc', '$'), ('Profit % of resale', 8, 'calc', '%')]
groups(ps, P0 - 2, [(1, 4, 'The house'), (5, 12, 'The two sales'), (13, 20, 'Calculated (rough: you don\'t know their real rehab)')])
header(ps, P0 - 1, pcols)
pf = {13: lambda r: f'=IF(OR(E{r}="",G{r}=""),"",(G{r}-E{r})/30.44)',
      14: lambda r: f'=IF(OR(F{r}="",H{r}=""),"",H{r}-F{r})',
      15: lambda r: f'=IF(OR(F{r}="",D{r}="",D{r}=0),"",F{r}/D{r})',
      16: lambda r: f'=IF(OR(H{r}="",D{r}="",D{r}=0),"",H{r}/D{r})',
      17: lambda r: f'=IF(H{r}="","",IF(I{r}<>"",I{r},IF(D{r}="","",D{r}*$E$3)))',
      18: lambda r: f'=IF(H{r}="","",H{r}*$E$4)',
      19: lambda r: f'=IF(OR(F{r}="",H{r}="",Q{r}=""),"",H{r}-F{r}-Q{r}-R{r})',
      20: lambda r: f'=IF(OR(S{r}="",H{r}="",H{r}=0),"",S{r}/H{r})'}
table(ps, P0, P1, pcols, formulas=pf, example={1: 'Sampleville, OH', 2: 'EX001', 3: '456 Sample Ave', 4: 1300, 5: dt.date(2025, 11, 10), 6: 118000,
                                               7: dt.date(2026, 5, 20), 8: 214900, 10: 'ABC Homes LLC', 11: 'Zillow sale history', 12: 'EXAMPLE row: delete me'})
ps.freeze_panes = ps.cell(row=P0, column=4)

# ============================ Calls ============================
N_C = 300; C0 = 8; C1 = C0 + N_C - 1
cs = wb.create_sheet('Calls')
title(cs, 'Calls', 'Every local you talk to, with the same questions (list on the Start Here tab). When three strangers say the same thing, you can trust it.')
ccols = [('Metro', 18, 'in', 't'), ('Date', 11, 'in', 'd'), ('Name', 16, 'in', 't'), ('Role', 15, 'in', 't'), ('Phone / email', 20, 'in', 't'), ('Company', 16, 'in', 't'),
         ('Renovated 3-bed sells for ($)', 12, 'in', '$'), ('Cosmetic rehab ($/sqft)', 10, 'in', '$'), ('Full rehab ($/sqft)', 10, 'in', '$'), ('Rent, 3-bed ($/mo)', 10, 'in', '$'),
         ('Investor competition (1 few – 5 crowded)', 11, 'in', 'n'), ('Best ZIPs / areas', 20, 'in', 't'), ('Avoid (and why)', 20, 'in', 't'),
         ('Works with out-of-state buyers? (Y/N)', 10, 'in', 't'), ('Referred you to', 16, 'in', 't'), ('Notes', 26, 'in', 't')]
groups(cs, C0 - 2, [(1, 6, 'Who'), (7, 11, 'Their numbers'), (12, 16, 'What they said')])
header(cs, C0 - 1, ccols)
table(cs, C0, C1, ccols, example={1: 'Sampleville, OH', 2: dt.date(2026, 10, 12), 3: 'Jane Example', 4: 'Agent', 5: 'jane@example.com', 6: 'Example Realty',
                                  7: 215000, 8: 25, 9: 55, 10: 1450, 11: 3, 12: 'North side, EX001', 13: 'Flood zone by the river', 14: 'Y', 15: 'Bob (contractor)', 16: 'EXAMPLE row: delete me'})
dv = DataValidation(type='list', formula1='"Agent,Property manager,Contractor,Lender,Title company,Investor,Wholesaler,Inspector,Appraiser,Other"', allow_blank=True)
cs.add_data_validation(dv); dv.add(f'D{C0}:D{C1}')
dv = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=True); cs.add_data_validation(dv); dv.add(f'K{C0}:K{C1}')
dv = DataValidation(type='list', formula1='"Y,N"', allow_blank=True); cs.add_data_validation(dv); dv.add(f'N{C0}:N{C1}')
cs.freeze_panes = cs.cell(row=C0, column=4)

# ============================ Paper Flips ============================
N_X = 300; X0 = 8; X1 = X0 + N_X - 1
xs = wb.create_sheet('Paper Flips')
title(xs, 'Paper Flips', 'Practice with zero risk: run your numbers on real listings (Deal Check tab), log them here, then fill in what really happened.')
xs['A3'] = 'Paper flips logged'; xs['A3'].font = font(bold=True)
xs['D3'] = f'=COUNTA(C{X0}:C{X1})'
xs['A4'] = 'Average resale-value miss (+/-)'; xs['A4'].font = font(bold=True)
xs['D4'] = f'=IFERROR(AVERAGE(O{X0}:O{X1}),"")'
xs['A5'] = 'Resale estimates within 10%'; xs['A5'].font = font(bold=True)
xs['D5'] = f'=IF(COUNT(O{X0}:O{X1})=0,"",COUNTIF(O{X0}:O{X1},"<=0.1")&" of "&COUNT(O{X0}:O{X1}))'
for k in ('D3', 'D4', 'D5'):
    xs[k].fill = F_KEY; xs[k].font = font(bold=True); xs[k].border = BOX
xs['D3'].number_format = '0'; xs['D4'].number_format = '0.0%'
xs['E4'] = 'Aim for under 10% before you buy for real.'; xs['E4'].font = font('595959', italic=True)
xcols = [('Date', 11, 'in', 'd'), ('Metro', 18, 'in', 't'), ('Address', 26, 'in', 't'), ('Sqft', 8, 'in', 'n'), ('List price ($)', 11, 'in', '$'),
         ('Your max offer ($)', 11, 'in', '$'), ('Your resale estimate ($)', 12, 'in', '$'), ('Your rehab estimate ($)', 11, 'in', '$'),
         ('Actually sold for, as-is ($)', 12, 'in', '$'), ('Renovated resale or best comp found later ($)', 14, 'in', '$'), ('Notes', 26, 'in', 't'),
         ('Would your offer have won?', 10, 'calc', 't'), ('Your offer vs. sold price', 9, 'calc', '%'), ('Resale estimate miss', 9, 'calc', '%'), ('Miss (+/-)', 8, 'calc', '%')]
groups(xs, X0 - 2, [(1, 4, 'The listing'), (5, 8, 'Your numbers (before you know the answer)'), (9, 11, 'What really happened'), (12, 15, 'How close you were')])
header(xs, X0 - 1, xcols)
xf = {12: lambda r: f'=IF(OR(F{r}="",I{r}=""),"",IF(F{r}>=I{r},"Yes","No"))',
      13: lambda r: f'=IF(OR(F{r}="",I{r}="",I{r}=0),"",F{r}/I{r}-1)',
      14: lambda r: f'=IF(OR(G{r}="",J{r}="",J{r}=0),"",G{r}/J{r}-1)',
      15: lambda r: f'=IF(N{r}="","",ABS(N{r}))'}
table(xs, X0, X1, xcols, formulas=xf, example={1: dt.date(2026, 10, 5), 2: 'Sampleville, OH', 3: '789 Practice Rd', 4: 1200, 5: 135000, 6: 112000, 7: 205000, 8: 35000,
                                               9: 128500, 10: 212000, 11: 'EXAMPLE row: delete me'})
xs.freeze_panes = xs.cell(row=X0, column=4)

# ============================ Deal Check ============================
ds = wb.create_sheet('Deal Check')
ds.sheet_view.showGridLines = False
title(ds, 'Deal Check', 'One house at a time. Yellow cells are yours. The example numbers are filled in so you can see how it works; type over them.')
for col, w in zip('ABCDEFG', (40, 16, 12, 12, 13, 12, 44)):
    ds.column_dimensions[col].width = w
def lab(cell, text, bold=False):
    ds[cell] = text; ds[cell].font = font(bold=bold)
def inp(cell, v, fmt=None, comment=None):
    ds[cell] = v; ds[cell].fill = F_IN; ds[cell].font = font(BLUE); ds[cell].border = BOX
    if fmt: ds[cell].number_format = FMT.get(fmt, fmt)
    if comment: note(ds, cell, comment)
def calc(cell, f, fmt=None, key=False):
    ds[cell] = f; ds[cell].fill = F_KEY if key else F_CALC; ds[cell].font = font(bold=key); ds[cell].border = BOX
    if fmt: ds[cell].number_format = FMT.get(fmt, fmt)
def sec(cell, text):
    ds[cell] = text; ds[cell].font = font(NAVY, True, size=12)

sec('A4', 'The house')
lab('A5', 'Address'); inp('B5', '789 Practice Rd', 't')
lab('A6', 'Metro'); inp('B6', 'Sampleville, OH', 't')
lab('A7', 'Square feet'); inp('B7', 1200, 'n')
lab('A8', 'List / asking price ($)'); inp('B8', 135000, '$')

sec('A10', 'Comps: renovated houses that SOLD nearby (last 6 months, within ~1 mile, similar size)')
for i, h in enumerate(['Address', 'Sold price ($)', 'Sqft', 'Sold on', 'Distance (mi)', '$/sqft'], 1):
    c = ds.cell(row=11, column=i, value=h); c.font = font('FFFFFF', True); c.fill = F_HEAD; c.border = BOX; c.alignment = Alignment(wrap_text=True, horizontal='center')
comps = [('12 Comp Ln', 212000, 1250, dt.date(2026, 8, 14), 0.4), ('34 Comp Ln', 205000, 1180, dt.date(2026, 7, 2), 0.7), ('56 Comp Ln', 219500, 1320, dt.date(2026, 9, 21), 0.9)]
for k in range(5):
    r = 12 + k
    vals = comps[k] if k < len(comps) else (None,) * 5
    inp(f'A{r}', vals[0], 't'); inp(f'B{r}', vals[1], '$'); inp(f'C{r}', vals[2], 'n'); inp(f'D{r}', vals[3], 'd'); inp(f'E{r}', vals[4], 'n1')
    calc(f'F{r}', f'=IF(OR(B{r}="",C{r}="",C{r}=0),"",B{r}/C{r})', '$')
lab('A18', 'Lowest comp $/sqft'); calc('B18', '=IFERROR(MIN(F12:F16),"")', '$')
lab('A19', 'Middle comp $/sqft'); calc('B19', '=IFERROR(MEDIAN(F12:F16),"")', '$')
lab('A20', 'Resale value at the LOWEST comp (what you should use)'); calc('B20', '=IF(OR(B18="",B18=0,B7=""),"",B18*B7)', '$', key=True)
lab('A21', 'Resale value at the middle comp (for reference)'); calc('B21', '=IF(OR(B19="",B19=0,B7=""),"",B19*B7)', '$')

sec('A23', 'Your numbers')
rows = [
    ('A24', 'Resale value you\'ll use ($), blank = lowest comp', 'B24', None, '$', 'Leave blank to use the lowest-comp value above. Only type a higher number if an agent AND a lender agree.'),
    ('A26', 'Purchase price ($)', 'B26', 120000, '$', None),
    ('A27', 'Buying closing costs (% of price)', 'B27', 0.02, '%', 'Your assumption. Ask the title company for a real quote.'),
    ('A28', 'Rehab estimate ($)', 'B28', 35000, '$', 'Use the average of 3 contractor bids.'),
    ('A29', 'Rehab cushion for surprises (%)', 'B29', 0.20, '%', 'Extra on top of the rehab estimate. 20% is a safe starting cushion.'),
    ('A30', 'Months you\'ll own it (buy to sale closing)', 'B30', 6, 'n', None),
    ('A31', 'Holding costs per month ($): taxes, insurance, utilities', 'B31', 650, '$', 'Ask an insurance agent and check the county tax bill.'),
    ('A32', 'Selling costs (% of resale): agent fees + closing', 'B32', 0.08, '%', 'Your assumption. Ask a local agent.'),
    ('A33', 'Loan (% of purchase + rehab). 0% = all cash', 'B33', 0, '%', 'Family money as cash = 0%. Hard-money lenders often lend 85–90% of purchase and 100% of rehab.'),
    ('A34', 'Loan interest (per year)', 'B34', 0.11, '%', 'Only used if the loan is more than 0%.'),
    ('A35', 'Loan points (% of loan)', 'B35', 0.02, '%', None),
    ('A36', 'Profit you want (% of resale)', 'B36', 0.15, '%', 'Used for your max offer below.'),
]
for a, t, b, v, f, cm in rows:
    lab(a, t); inp(b, v, f, cm)
lab('A25', 'Resale value used ($)'); calc('B25', '=IF(B24<>"",B24,B20)', '$', key=True)

sec('A38', 'Result')
res = [
    ('A39', 'Buying closing costs ($)', 'B39', '=B26*B27', '$'),
    ('A40', 'Rehab with cushion ($)', 'B40', '=B28*(1+B29)', '$'),
    ('A41', 'Holding costs ($)', 'B41', '=B30*B31', '$'),
    ('A42', 'Loan amount ($)', 'B42', '=B33*(B26+B40)', '$'),
    ('A43', 'Loan interest + points ($)', 'B43', '=B42*B34*B30/12+B42*B35', '$'),
    ('A44', 'Selling costs ($)', 'B44', '=IF(B25="","",B25*B32)', '$'),
    ('A45', 'Total costs, everything ($)', 'B45', '=IF(B25="","",B26+B39+B40+B41+B43+B44)', '$'),
    ('A47', 'Cash you need to put in ($)', 'B47', '=B26+B39+B40+B41+B43-B42', '$'),
]
for a, t, b, f, fm in res:
    lab(a, t); calc(b, f, fm)
lab('A46', 'PROFIT ($)', True); calc('B46', '=IF(B25="","",B25-B45)', '$', key=True)
lab('A48', 'Profit as % of resale', True); calc('B48', '=IF(OR(B46="",B25=0),"",B46/B25)', '%', key=True)
lab('A49', 'Return on your cash', True); calc('B49', '=IF(OR(B46="",B47<=0),"",B46/B47)', '%', key=True)
lab('A50', 'Your max offer for the profit you want ($)', True)
calc('B50', '=IF(B25="","",(B25*(1-B32-B36)-B40-B41-B33*B40*(B34*B30/12+B35))/(1+B27+B33*(B34*B30/12+B35)))', '$', key=True)
ds['C50'] = 'Includes loan costs. Compare it to the asking price.'; ds['C50'].font = font('595959', italic=True)
lab('A51', '"70% rule" check: resale × 70% − rehab ($)'); calc('B51', '=IF(B25="","",B25*0.7-B40)', '$')
ds['C51'] = 'A common quick rule. Your own max offer above is more precise.'; ds['C51'].font = font('595959', italic=True)

sec('A53', 'Stress test: does it still make money if things go wrong?')
for i, h in enumerate(['What goes wrong', 'Resale ($)', 'Rehab ($)', 'Months', 'Profit ($)', 'Profit %'], 1):
    c = ds.cell(row=54, column=i, value=h); c.font = font('FFFFFF', True); c.fill = F_HEAD; c.border = BOX; c.alignment = Alignment(horizontal='center')
sc = [('As planned', '=B25', '=B28', '=B30'),
      ('Resale 10% lower', '=B25*0.9', '=B28', '=B30'),
      ('Rehab 20% higher', '=B25', '=B28*1.2', '=B30'),
      ('Takes 3 months longer', '=B25', '=B28', '=B30+3'),
      ('All three at once', '=B25*0.9', '=B28*1.2', '=B30+3')]
for k, (name, a, rh, m) in enumerate(sc):
    r = 55 + k
    ds.cell(row=r, column=1, value=name).font = font(bold=(k == 4)); ds.cell(row=r, column=1).border = BOX
    calc(f'B{r}', f'=IF($B$25="","",{a[1:]})', '$'); calc(f'C{r}', rh, '$'); calc(f'D{r}', m, 'n')
    prof = (f'=IF(B{r}="","",B{r}*(1-$B$32)-$B$26*(1+$B$27)-C{r}*(1+$B$29)-D{r}*$B$31'
            f'-$B$33*($B$26+C{r}*(1+$B$29))*($B$34*D{r}/12+$B$35))')
    calc(f'E{r}', prof, '$', key=(k == 4)); calc(f'F{r}', f'=IF(OR(E{r}="",B{r}=0),"",E{r}/B{r})', '%')
lab('A61', 'Passes the stress test?', True)
calc('B61', '=IF(E59="","",IF(E59>0,"YES: still profitable","NO: too thin"))', None, key=True)
ds.merge_cells('B61:D61')

sec('A63', 'Second opinions (before you sign anything)')
chk = [('A64', 'A local agent agrees with the resale value (name)', 'B64', None), ('A65', 'A lender would lend on this deal (Y/N)', 'B65', None),
       ('A66', 'Contractor bids you got (how many)', 'B66', None), ('A67', 'Inspection done, no big surprises (Y/N)', 'B67', None),
       ('A68', 'Logged in Paper Flips tab (Y/N)', 'B68', None)]
for a, t, b, v in chk:
    lab(a, t); inp(b, v, 't')
ds.merge_cells('B64:D64')
ds.freeze_panes = 'A4'

# ============================ Flip Candidates ============================
F_AUTO = PatternFill('solid', fgColor='DDEBF7')
N_F = 200; FC0 = 8; FC1 = FC0 + N_F - 1
fs = wb.create_sheet('Flip Candidates')
title(fs, 'Flip Candidates', 'Likely renovated resales: older houses that sold well above their ZIP\'s typical $/sqft. Open each link, look at the sale history, and type the EARLIER sale (yellow).')
fcols = [('Metro', 18, 'in', 't'), ('ZIP', 8, 'in', 't'), ('Address', 26, 'in', 't'), ('Link', 12, 'in', 't'), ('Sqft', 8, 'in', 'n'), ('Year built', 7, 'in', 'yr'),
         ('Sold on', 11, 'in', 'd'), ('Sold for ($)', 11, 'in', '$'), ('ZIP typical $/sqft', 9, 'in', '$'),
         ('Earlier sale date (from history)', 12, 'in', 'd'), ('Earlier sale price ($)', 12, 'in', '$'), ('Notes', 22, 'in', 't'),
         ('Sold $/sqft', 8, 'calc', '$'), ('Above ZIP $/sqft', 8, 'calc', '%'), ('Flip?', 10, 'calc', 't'), ('Months between sales', 8, 'calc', 'n1'),
         ('Est. profit ($)', 11, 'calc', '$'), ('Profit % of resale', 8, 'calc', '%')]
groups(fs, FC0 - 2, [(1, 9, 'From the Redfin "sold" download (OfferDesk fills these)'), (10, 12, 'You check the sale history'), (13, 18, 'Calculated')])
header(fs, FC0 - 1, fcols)
ff = {13: lambda r: f'=IF(OR(H{r}="",E{r}="",E{r}=0),"",H{r}/E{r})',
      14: lambda r: f'=IF(OR(M{r}="",I{r}="",I{r}=0),"",M{r}/I{r}-1)',
      15: lambda r: f'=IF(OR(J{r}="",K{r}="",G{r}="",K{r}=0),"",IF(AND(G{r}-J{r}<=550,H{r}/K{r}-1>=0.25),"Likely flip","Not a flip"))',
      16: lambda r: f'=IF(OR(J{r}="",G{r}=""),"",(G{r}-J{r})/30.44)',
      17: lambda r: f'=IF(O{r}<>"Likely flip","",H{r}-K{r}-E{r}*{PF}!$E$3-H{r}*{PF}!$E$4)',
      18: lambda r: f'=IF(OR(Q{r}="",H{r}=0),"",Q{r}/H{r})'}
table(fs, FC0, FC1, fcols, formulas=ff, example={1: 'Sampleville, OH', 2: 'EX001', 3: '22 Flip St', 5: 1250, 6: 1952, 7: dt.date(2026, 8, 30), 8: 209000, 9: 142,
                                                  10: dt.date(2025, 12, 5), 11: 112000, 12: 'EXAMPLE row: delete me'})
fs.freeze_panes = fs.cell(row=FC0, column=4)
note(fs, 'O7', 'Likely flip = the earlier sale was within ~18 months and the price jumped 25% or more. Profit uses the rehab and cost assumptions at the top of the Past Flips tab.')

# ============================ Competition ============================
N_K = 200; K0 = 8; K1 = K0 + N_K - 1
ks = wb.create_sheet('Competition')
title(ks, 'Competition', 'Companies that advertise buying houses for cash in each metro ("we buy houses"). Maya can fill this from a web search; add any you find yourself.')
kcols = [('Metro', 18, 'in', 't'), ('Company', 26, 'in', 't'), ('Website', 30, 'in', 't'), ('Found by', 10, 'in', 't'), ('Local? (Y/N)', 8, 'in', 't'), ('Notes', 30, 'in', 't')]
header(ks, K0 - 1, kcols)
table(ks, K0, K1, kcols, example={1: 'Sampleville, OH', 2: 'Example Home Buyers', 3: 'https://example.com', 4: 'You', 5: 'Y', 6: 'EXAMPLE row: delete me'})
dv = DataValidation(type='list', formula1='"Y,N"', allow_blank=True); ks.add_data_validation(dv); dv.add(f'E{K0}:E{K1}')
ks.freeze_panes = ks.cell(row=K0, column=2)

# ============================ Metros ============================
ms = wb.create_sheet('Metros', 1)
title(ms, 'Metros', 'Your candidate metros side by side. Most columns fill themselves from the other tabs. Set your own weights in row 6; the score compares your metros to each other.')
M0, M1 = 8, 13
Z = lambda col: f"ZIPs!${col}${Z0}:${col}${Z1}"
L = lambda col: f"{LW}!${col}${L0}:${col}${L1}"
P = lambda col: f"{PF}!${col}${P0}:${col}${P1}"
C = lambda col: f"Calls!${col}${C0}:${col}${C1}"
mcols = [('Metro (exact same name in every tab)', 20, 'in', 't'),
         ('ZIPs researched', 8, 'link', 'n'), ('Typical median price ($)', 11, 'link', '$'), ('Sales, last 3 mo (your ZIPs)', 9, 'link', 'n'),
         ('Avg days on market', 8, 'link', 'n'), ('Avg price change, 1 yr', 8, 'link', '%'), ('Avg rent yield', 8, 'link', '%'), ('Avg share of houses rented out', 9, 'link', '%'),
         ('Flips found, 12 mo', 8, 'link', 'n'), ('Flips as % of sales (est.)', 9, 'calc', '%'), ('Avg est. flip profit (% of resale)', 9, 'link', '%'),
         ('Listings watched', 8, 'link', 'n'), ('Avg days to pending', 8, 'link', 'n'), ('Share with a price cut', 8, 'link', '%'),
         ('Calls logged', 7, 'link', 'n'), ('Avg competition from calls (1–5)', 9, 'link', 'n1'), ('Avg renovated 3-bed price, calls ($)', 11, 'link', '$'), ('Avg cosmetic rehab $/sqft, calls', 9, 'link', '$'),
         ('"We buy houses" companies (Competition tab)', 9, 'link', 'n'), ('County jobs, 2 yrs', 8, 'in', '%'), ('County population since 2020', 9, 'in', '%'), ('Your rating after visiting (1–5)', 9, 'in', 'n'),
         ('YOUR SCORE (0–100)', 10, 'calc', '0')]
FACT = [('Deal volume', 'D', 0), ('Sells fast', 'E', 1), ('Price trend', 'F', 0), ('Rent yield', 'G', 0), ('Fewer flippers', 'J', 1),
        ('Calls say less crowded', 'P', 1), ('Fewer "we buy houses"', 'S', 1), ('Job growth', 'T', 0), ('Population growth', 'U', 0), ('Your rating', 'V', 0)]
WTS = [15, 10, 10, 5, 15, 15, 10, 8, 7, 5]
for name, col, low in FACT:
    mcols.append((f'Points: {name}', 8, 'calc', '0.00'))
groups(ms, M0 - 2, [(1, 1, ''), (2, 8, 'From your ZIPs tab'), (9, 11, 'From Past Flips + Flip Candidates'), (12, 14, 'From Listing Watch'), (15, 18, 'From Calls'), (19, 19, 'Competition'), (20, 22, 'You type'), (23, 23, 'Result')])
header(ms, M0 - 1, mcols)
# weights row sits above the points columns, in the group row (row 6)
P0C = 24
for k, w in enumerate(WTS):
    c = ms.cell(row=M0 - 2, column=P0C + k, value=w)
    c.fill = F_IN; c.font = font(BLUE, True); c.border = BOX; c.alignment = Alignment(horizontal='center'); c.number_format = '0'
ms.cell(row=M0 - 3, column=P0C, value='Your weights (row 6, should add to 100) →').font = font(NAVY, True)
ms.cell(row=M0 - 3, column=P0C + len(WTS), value=f'=SUM({CL(P0C)}{M0-2}:{CL(P0C+len(WTS)-1)}{M0-2})').font = font(bold=True)
ms.cell(row=M0 - 3, column=P0C + len(WTS) + 1, value='total').font = font('595959', italic=True)
def nz(r, f):
    return f'=IF($A{r}="","",IF($B{r}=0,"",{f}))'
mf = {
    2: lambda r: f'=IF($A{r}="","",COUNTIF({Z("A")},$A{r}))',
    3: lambda r: nz(r, f'IFERROR(AVERAGEIFS({Z("D")},{Z("A")},$A{r}),"")'),
    4: lambda r: nz(r, f'SUMIFS({Z("E")},{Z("A")},$A{r})'),
    5: lambda r: nz(r, f'IFERROR(AVERAGEIFS({Z("G")},{Z("A")},$A{r}),"")'),
    6: lambda r: nz(r, f'IFERROR(AVERAGEIFS({Z("J")},{Z("A")},$A{r}),"")'),
    7: lambda r: nz(r, f'IFERROR(AVERAGEIFS({Z("S")},{Z("A")},$A{r}),"")'),
    8: lambda r: nz(r, f'IFERROR(AVERAGEIFS({Z("T")},{Z("A")},$A{r}),"")'),
    9: lambda r: f'=IF($A{r}="","",COUNTIF({P("A")},$A{r})+COUNTIFS({FC}!$A${FC0}:$A${FC1},$A{r},{FC}!$O${FC0}:$O${FC1},"Likely flip"))',
    19: lambda r: f'=IF($A{r}="","",IF(COUNTIF(Competition!$A${K0}:$A${K1},$A{r})=0,"",COUNTIF(Competition!$A${K0}:$A${K1},$A{r})))',
    10: lambda r: f'=IF(OR($A{r}="",D{r}="",D{r}=0),"",I{r}/(D{r}*4))',
    11: lambda r: (f'=IF($A{r}="","",IFERROR((SUMIFS({P("T")},{P("A")},$A{r})+SUMIFS({FC}!$R${FC0}:$R${FC1},{FC}!$A${FC0}:$A${FC1},$A{r},{FC}!$O${FC0}:$O${FC1},"Likely flip"))'
                   f'/(COUNTIFS({P("A")},$A{r},{P("T")},"<>")+COUNTIFS({FC}!$A${FC0}:$A${FC1},$A{r},{FC}!$O${FC0}:$O${FC1},"Likely flip")),""))'),
    12: lambda r: f'=IF($A{r}="","",COUNTIF({L("A")},$A{r}))',
    13: lambda r: f'=IF($A{r}="","",IFERROR(AVERAGEIFS({L("S")},{L("A")},$A{r}),""))',
    14: lambda r: f'=IF(OR($A{r}="",L{r}=0),"",COUNTIFS({L("A")},$A{r},{L("R")},">0")/L{r})',
    15: lambda r: f'=IF($A{r}="","",COUNTIF({C("A")},$A{r}))',
    16: lambda r: f'=IF($A{r}="","",IFERROR(AVERAGEIFS({C("K")},{C("A")},$A{r}),""))',
    17: lambda r: f'=IF($A{r}="","",IFERROR(AVERAGEIFS({C("G")},{C("A")},$A{r}),""))',
    18: lambda r: f'=IF($A{r}="","",IFERROR(AVERAGEIFS({C("H")},{C("A")},$A{r}),""))',
}
first_p, last_p = CL(P0C), CL(P0C + len(FACT) - 1)
wr = f'${first_p}${M0-2}:${last_p}${M0-2}'
mf[23] = lambda r: (f'=IF($A{r}="","",IFERROR(ROUND(100*SUMPRODUCT({wr},{first_p}{r}:{last_p}{r})/'
                    f'SUMPRODUCT({wr},--({first_p}{r}:{last_p}{r}<>"")),0),""))')
for k, (name, col, low) in enumerate(FACT):
    rng = f'{col}${M0}:{col}${M1}'
    mf[P0C + k] = (lambda col=col, low=low, rng=rng: (lambda r: f'=IF(OR($A{r}="",{col}{r}=""),"",IF(COUNT({rng})<2,1,(COUNT({rng})-RANK({col}{r},{rng},{low}))/(COUNT({rng})-1)))'))()
table(ms, M0, M1, mcols, formulas=mf, example={1: 'Sampleville, OH', 20: 0.018, 21: 0.024})
ms.cell(row=M0, column=1).font = font(BLUE, italic=True)
ms.freeze_panes = ms.cell(row=M0, column=2)
ms.conditional_formatting.add(f'W{M0}:W{M1}', ColorScaleRule(start_type='num', start_value=0, start_color='F8696B', mid_type='num', mid_value=50, mid_color='FFEB84', end_type='num', end_value=100, end_color='63BE7B'))
dv = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=True); ms.add_data_validation(dv); dv.add(f'V{M0}:V{M1}')
note(ms, 'W7', 'For each factor, your metros are ranked against each other: best = 1 point, worst = 0. The score is the weighted average of the points it has data for, times 100. It only compares the metros in this table.')
note(ms, 'H7', 'Average share of single-family homes that are rented out (mostly investor-owned). Higher = more investors already here.')
note(ms, 'J7', 'Flips you found ÷ yearly sales in your ZIPs. Only meaningful after you have searched sale histories properly.')

# Ready check
R0 = M1 + 3
ms.cell(row=R0, column=1, value='Ready check: don\'t buy in a metro until every box says Yes').font = font(NAVY, True, size=12)
rcols = [('Metro', 20), ('Data spot-checked vs. Redfin (Y/N)', 11), ('10+ local calls', 11), ('3+ agents called', 11), ('A lender reviewed a real deal (Y/N)', 11),
         ('10+ paper flips', 11), ('Paper resale miss ≤ 10%', 11), ('Visited in person (Y/N)', 11), ('READY?', 11)]
for i, (h, w) in enumerate(rcols, 1):
    c = ms.cell(row=R0 + 1, column=i, value=h); c.font = font('FFFFFF', True); c.fill = F_HEAD; c.border = BOX
    c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
ms.row_dimensions[R0 + 1].height = 42
XF = "'Paper Flips'"
for k in range(M1 - M0 + 1):
    r = R0 + 2 + k; mr = M0 + k
    def setc(col, val, kind):
        c = ms.cell(row=r, column=col, value=val); c.border = BOX
        if kind == 'in':
            c.fill = F_IN; c.font = font(BLUE)
        else:
            c.fill = F_CALC; c.font = font(GREEN if kind == 'link' else BLACK, bold=(col == 9))
        c.alignment = Alignment(horizontal='center') if col > 1 else Alignment()
    setc(1, f'=IF($A${mr}="","",$A${mr})', 'calc')
    setc(2, None, 'in'); setc(5, None, 'in'); setc(8, None, 'in')
    setc(3, f'=IF($A{r}="","",IF(COUNTIF({C("A")},$A{r})>=10,"Yes","No ("&COUNTIF({C("A")},$A{r})&"/10)"))', 'link')
    setc(4, f'=IF($A{r}="","",IF(COUNTIFS({C("A")},$A{r},{C("D")},"Agent")>=3,"Yes","No ("&COUNTIFS({C("A")},$A{r},{C("D")},"Agent")&"/3)"))', 'link')
    setc(6, f'=IF($A{r}="","",IF(COUNTIF({XF}!$B${X0}:$B${X1},$A{r})>=10,"Yes","No ("&COUNTIF({XF}!$B${X0}:$B${X1},$A{r})&"/10)"))', 'link')
    setc(7, f'=IF($A{r}="","",IF(COUNTIFS({XF}!$B${X0}:$B${X1},$A{r},{XF}!$O${X0}:$O${X1},"<>")=0,"No data yet",IF(AVERAGEIFS({XF}!$O${X0}:$O${X1},{XF}!$B${X0}:$B${X1},$A{r})<=0.1,"Yes","No ("&TEXT(AVERAGEIFS({XF}!$O${X0}:$O${X1},{XF}!$B${X0}:$B${X1},$A{r}),"0%")&")")))', 'link')
    setc(9, f'=IF($A{r}="","",IF(AND(UPPER(B{r})="Y",C{r}="Yes",D{r}="Yes",UPPER(E{r})="Y",F{r}="Yes",G{r}="Yes",UPPER(H{r})="Y"),"READY","Not yet"))', 'calc')
rr = f'I{R0+2}:I{R0+2+M1-M0}'
ms.conditional_formatting.add(rr, CellIsRule(operator='equal', formula=['"READY"'], fill=PatternFill('solid', fgColor='C6EFCE'), font=Font(name=ARIAL, color='006100', bold=True)))
ms.conditional_formatting.add(rr, CellIsRule(operator='equal', formula=['"Not yet"'], fill=PatternFill('solid', fgColor='FFC7CE'), font=Font(name=ARIAL, color='9C0006', bold=True)))
yn = DataValidation(type='list', formula1='"Y,N"', allow_blank=True); ms.add_data_validation(yn)
for col in 'BEH':
    yn.add(f'{col}{R0+2}:{col}{R0+2+M1-M0}')
ms.cell(row=R0 + 2 + (M1 - M0) + 2, column=1, value='Paper resale miss uses the "Miss (+/-)" column in Paper Flips, which needs a renovated resale or comp you found later.').font = font('595959', italic=True)

# order tabs
order = ['Start Here', 'Metros', 'ZIPs', 'Listing Watch', 'Flip Candidates', 'Past Flips', 'Competition', 'Calls', 'Deal Check', 'Paper Flips']
wb._sheets = [wb[n] for n in order]
for w in wb.worksheets:
    w.sheet_properties.tabColor = {'Start Here': NAVY, 'Metros': '548235'}.get(w.title, '8EA9DB')
wb.save(OUT)
print('saved', OUT)
